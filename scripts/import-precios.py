"""Genera lib/catalog-seed.ts desde el Excel oficial de Genera."""
from __future__ import annotations

import re
from pathlib import Path

import openpyxl

EXCEL = Path(r"C:\Users\Admin\Desktop\PRECIOS DE EQUIPOS PARA IA.xlsx")
OUT = Path(r"C:\Users\Admin\solar-financials\lib\catalog-seed.ts")
TRM = 4125
NOW = "2026-10-10T00:00:00.000Z"

HEADER_WORDS = {"CATEGORY", "MARCA", "PRECIO_COP", "COST_COP"}

GROUPS = [
    ("A", "B", "C"),
    ("E", "F", "G"),
    ("I", "J", "K"),
]


def slug(text: str) -> str:
    s = text.lower().strip()
    s = re.sub(r"[^a-z0-9]+", "-", s)
    return s.strip("-")[:48] or "item"


def parse_kw(text: str) -> float | None:
    m = re.search(r"(\d+(?:[.,]\d+)?)\s*k[wW]", text)
    if m:
        return float(m.group(1).replace(",", "."))
    m = re.search(r"^(\d+)\s*W\b", text, re.I)
    if m:
        return float(m.group(1))
    return None


def map_category(section: str, name: str) -> tuple[str, str | None, str, str | None]:
    sec = re.sub(r"\s+", " ", section.upper().strip())
    n = name.upper()
    notes = None
    unit = "unidad"
    mount = None
    cat = "cabling"

    if "PANEL" in sec:
        cat, unit = "panel", "módulo"
    elif sec.startswith("INVERSOR"):
        cat = "inverter"
    elif "BATER" in sec:
        cat, unit = "battery", "módulo"
    elif "ESTRUCTURA" in sec:
        cat, mount = "structure", "cubierta"
        if "STACK" in n:
            mount = "suelo"
    elif "PROTECCIONES DC" in sec or sec == "DPS" or "COMBINER" in sec:
        cat, notes = "protection", "DC"
    elif "PROTECCIONES AC" in sec or "CAJA PROTECC" in sec:
        cat, notes = "protection", "AC"
    elif "VARILLA" in sec:
        cat = "protection"
    elif "MICROINVERSOR" in sec:
        if any(k in n for k in ("DS3", "QT2", "HMS", "DEYE 2000", "DEYE 1600")):
            cat, notes = "inverter", "Microinversor"
        else:
            cat, notes = "cabling", "Monitoreo / accesorio"
    elif "CARGADOR" in sec:
        cat, notes, unit = "cabling", "Cargador EV", "unidad"
    elif "CABLE" in sec or "CORAZA" in sec or "CANALETA" in sec or sec == "MC4":
        cat = "cabling"
        unit = "metro" if any(k in n for k in ("MM", "CABLE", "CORAZA", "ELECTROSTATICA", "CANALETA")) else "unidad"
    else:
        cat = "cabling"

    return cat, notes, unit, mount


def ts_escape(s: str) -> str:
    return s.replace("\\", "\\\\").replace('"', '\\"')


def main() -> None:
    wb = openpyxl.load_workbook(EXCEL, data_only=True)
    ws = wb.active
    items: list[dict] = []
    seen: set[str] = set()
    section = {"A": "", "E": "", "I": ""}

    for r in range(1, ws.max_row + 1):
        for name_col, brand_col, price_col in GROUPS:
            name = ws[f"{name_col}{r}"].value
            brand = ws[f"{brand_col}{r}"].value
            price = ws[f"{price_col}{r}"].value
            if name is None:
                continue
            name_s = str(name).strip()
            if not name_s:
                continue
            if brand is None and price is None:
                section[name_col] = name_s
                continue
            if str(name_s).upper() in HEADER_WORDS:
                continue
            if not isinstance(price, (int, float)):
                continue
            brand_s = str(brand).strip() if brand else "Genera"
            cat, notes, unit, mount = map_category(section[name_col], name_s)
            ident = f"{cat[:4]}-{slug(brand_s)}-{slug(name_s)}"
            base = ident
            n = 2
            while ident in seen:
                ident = f"{base}-{n}"
                n += 1
            seen.add(ident)

            cap = parse_kw(name_s)
            is_default = False
            if ident == "pane-znshine-solar-650w":
                is_default = True
            if "12kw" in name_s.lower() and "trifasico lv" in name_s.lower().replace(" ", "") and brand_s.lower() == "deye":
                is_default = True
            if name_s.lower().replace(" ", "") == "10kw-lv" and "felicity" in brand_s.lower():
                is_default = True
            if name_s.lower().startswith("riel 4.80") and "improinde" in brand_s.lower() and "m" not in name_s.lower()[4:]:
                is_default = True
            if name_s.lower() == "6mm" and "kibor" in brand_s.lower():
                is_default = True
            if "20a" in name_s.lower() and "500v" in name_s.lower() and brand_s.upper() == "YRO":
                is_default = True

            items.append(
                {
                    "id": ident,
                    "category": cat,
                    "name": name_s,
                    "brand": re.sub(r"\s+", " ", brand_s),
                    "model": name_s,
                    "priceCOP": int(round(price)),
                    "unit": unit,
                    "capacity": cap,
                    "mountType": mount,
                    "isDefault": is_default,
                    "notes": notes,
                }
            )

    # Ensure at least one default per core category
    defaults_needed = {
        "panel": lambda i: "650" in i["name"],
        "inverter": lambda i: "12KW TRIFASICO LV" in i["name"].upper().replace("  ", " ") and i["brand"] == "Deye",
        "battery": lambda i: i["name"].upper().replace(" ", "") == "10KW-LV" and "Felicity" in i["brand"],
        "structure": lambda i: i["name"].lower().startswith("riel 4.80") and "Improinde" in i["brand"],
        "cabling": lambda i: i["name"] == "6mm" and i["brand"] == "Kibor",
        "protection": lambda i: "20A" in i["name"] and "YRO" in i["brand"],
    }
    have = {c: False for c in defaults_needed}
    for i in items:
        if i["isDefault"]:
            have[i["category"]] = True
    for cat, pred in defaults_needed.items():
        if have.get(cat):
            continue
        for i in items:
            if i["category"] == cat and pred(i):
                i["isDefault"] = True
                break

    lines = [
        'import type { EquipmentCategory, InventoryItem, StructureMountType } from "@/types/inventory";',
        "",
        "const TRM = 4125;",
        f'const NOW = "{NOW}";',
        "",
        "function item(input: {",
        "  id: string;",
        "  category: EquipmentCategory;",
        "  name: string;",
        "  brand: string;",
        "  model: string;",
        "  priceCOP: number;",
        "  unit?: string;",
        "  capacity?: number;",
        "  mountType?: StructureMountType;",
        "  isDefault?: boolean;",
        "  notes?: string;",
        "}): InventoryItem {",
        "  return {",
        '    unit: "unidad",',
        "    isDefault: false,",
        "    active: true,",
        "    updatedAt: NOW,",
        "    ...input,",
        "    priceUSD: Number((input.priceCOP / TRM).toFixed(2)),",
        "  };",
        "}",
        "",
        "/** Catálogo oficial importado de PRECIOS DE EQUIPOS PARA IA.xlsx */",
        "export const SEED_INVENTORY: InventoryItem[] = [",
    ]
    for i in items:
        fields = [
            f'id: "{i["id"]}"',
            f'category: "{i["category"]}"',
            f'name: "{ts_escape(i["name"])}"',
            f'brand: "{ts_escape(i["brand"])}"',
            f'model: "{ts_escape(i["model"])}"',
            f'priceCOP: {i["priceCOP"]}',
        ]
        if i["unit"] != "unidad":
            fields.append(f'unit: "{i["unit"]}"')
        if i["capacity"] is not None:
            fields.append(f'capacity: {i["capacity"]}')
        if i["mountType"]:
            fields.append(f'mountType: "{i["mountType"]}"')
        if i["isDefault"]:
            fields.append("isDefault: true")
        if i["notes"]:
            fields.append(f'notes: "{ts_escape(i["notes"])}"')
        lines.append("  item({")
        lines.append("    " + ",\n    ".join(fields) + ",")
        lines.append("  }),")
    lines.append("];")
    lines.append("")
    OUT.write_text("\n".join(lines), encoding="utf-8")
    print(f"Wrote {len(items)} items to {OUT}")
    for cat in ("panel", "inverter", "battery", "structure", "cabling", "protection"):
        d = [i for i in items if i["category"] == cat and i["isDefault"]]
        print(cat, "count", sum(1 for i in items if i["category"] == cat), "default", d[0]["name"] if d else "NONE")


if __name__ == "__main__":
    main()
